# Django Management Command for Automated Report Generation
import os
from datetime import datetime, timedelta
from django.core.management.base import BaseCommand
from django.conf import settings
from django.utils import timezone
from django.core.mail import EmailMessage

from riders.reports import ReportGenerator


class Command(BaseCommand):
    help = 'Generate and email automated reports'

    def add_arguments(self, parser):
        parser.add_argument(
            '--report-type',
            type=str,
            default='attendance',
            choices=['attendance', 'rider_summary', 'module_completion', 'geographical', 'all'],
            help='Type of report to generate'
        )
        parser.add_argument(
            '--format',
            type=str,
            default='excel',
            choices=['csv', 'excel', 'both'],
            help='Export format'
        )
        parser.add_argument(
            '--days-back',
            type=int,
            default=7,
            help='Number of days back to include in report'
        )
        parser.add_argument(
            '--email',
            type=str,
            help='Email address to send report to'
        )
        parser.add_argument(
            '--save-path',
            type=str,
            default='/tmp/digitalboda_reports',
            help='Path to save reports'
        )

    def handle(self, *args, **options):
        self.stdout.write(self.style.SUCCESS('Starting automated report generation...'))
        
        try:
            # Setup filters for the report
            end_date = timezone.now().date()
            start_date = end_date - timedelta(days=options['days_back'])
            
            filters = {
                'start_date': start_date,
                'end_date': end_date
            }
            
            # Create save directory
            save_path = options['save_path']
            os.makedirs(save_path, exist_ok=True)
            
            # Initialize report generator
            generator = ReportGenerator()
            timestamp = datetime.now().strftime('%Y%m%d_%H%M')
            
            generated_files = []
            
            # Generate reports based on type
            if options['report_type'] in ['attendance', 'all']:
                generated_files.extend(
                    self._generate_report(generator, 'attendance', timestamp, save_path, options, filters)
                )
            
            if options['report_type'] in ['rider_summary', 'all']:
                generated_files.extend(
                    self._generate_report(generator, 'rider_summary', timestamp, save_path, options, filters)
                )
            
            if options['report_type'] in ['module_completion', 'all']:
                generated_files.extend(
                    self._generate_report(generator, 'module_completion', timestamp, save_path, options, filters)
                )
            
            if options['report_type'] in ['geographical', 'all']:
                generated_files.extend(
                    self._generate_report(generator, 'geographical', timestamp, save_path, options, filters)
                )
            
            # Send email if requested
            if options['email'] and generated_files:
                self._send_email_report(options['email'], generated_files, filters)
            
            self.stdout.write(
                self.style.SUCCESS(
                    f'Successfully generated {len(generated_files)} report files'
                )
            )
            
            for file_path in generated_files:
                self.stdout.write(f'  - {file_path}')
        
        except Exception as e:
            self.stdout.write(
                self.style.ERROR(f'Error generating reports: {str(e)}')
            )

    def _generate_report(self, generator, report_type, timestamp, save_path, options, filters):
        """Generate a specific report type"""
        generated_files = []
        
        try:
            # Get the data
            if report_type == 'attendance':
                data = generator.generate_attendance_report(**filters)
            elif report_type == 'rider_summary':
                data = generator.generate_rider_summary_report(**filters)
            elif report_type == 'module_completion':
                data = generator.generate_module_completion_report(**filters)
            elif report_type == 'geographical':
                data = generator.generate_geographical_report(**filters)
            else:
                return generated_files
            
            if not data:
                self.stdout.write(
                    self.style.WARNING(f'No data found for {report_type} report')
                )
                return generated_files
            
            # Generate CSV if requested
            if options['format'] in ['csv', 'both']:
                csv_filename = f"{report_type}_report_{timestamp}.csv"
                csv_path = os.path.join(save_path, csv_filename)
                self._save_csv(data, csv_path)
                generated_files.append(csv_path)
            
            # Generate Excel if requested
            if options['format'] in ['excel', 'both']:
                excel_filename = f"{report_type}_report_{timestamp}.xlsx"
                excel_path = os.path.join(save_path, excel_filename)
                self._save_excel(data, excel_path, report_type.replace('_', ' ').title())
                generated_files.append(excel_path)
            
            self.stdout.write(
                self.style.SUCCESS(f'Generated {report_type} report with {len(data)} records')
            )
        
        except Exception as e:
            self.stdout.write(
                self.style.ERROR(f'Error generating {report_type} report: {str(e)}')
            )
        
        return generated_files

    def _save_csv(self, data, file_path):
        """Save data as CSV file"""
        import csv
        
        with open(file_path, 'w', newline='', encoding='utf-8') as csvfile:
            if data:
                writer = csv.DictWriter(csvfile, fieldnames=data[0].keys())
                writer.writeheader()
                writer.writerows(data)

    def _save_excel(self, data, file_path, sheet_name):
        """Save data as Excel file"""
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        
        if not data:
            return
        
        # Create DataFrame
        df = pd.DataFrame(data)
        
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        # Add title
        ws['A1'] = f"DigitalBoda {sheet_name} - Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:E1')
        
        # Add headers starting from row 3
        headers = list(df.columns)
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # Add data
        for row_num, row_data in enumerate(df.values, 4):
            for col_num, value in enumerate(row_data, 1):
                ws.cell(row=row_num, column=col_num, value=value)
        
        # Auto-adjust column widths
        for col_num in range(1, len(headers) + 1):
            max_length = 0
            column_letter = get_column_letter(col_num)
            
            # Check header length
            header_length = len(str(headers[col_num - 1]))
            if header_length > max_length:
                max_length = header_length
            
            # Check data lengths
            for row_num in range(4, ws.max_row + 1):
                cell_value = ws.cell(row=row_num, column=col_num).value
                if cell_value:
                    cell_length = len(str(cell_value))
                    if cell_length > max_length:
                        max_length = cell_length
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(file_path)

    def _send_email_report(self, recipient, file_paths, filters):
        """Send generated reports via email"""
        try:
            subject = f"DigitalBoda Automated Report - {datetime.now().strftime('%Y-%m-%d')}"
            
            message = f"""
            Dear Team,

            Please find attached the automated DigitalBoda reports generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}.

            Report Period: {filters.get('start_date', 'N/A')} to {filters.get('end_date', 'N/A')}
            Files Generated: {len(file_paths)}

            Best regards,
            DigitalBoda Reporting System
            """
            
            email = EmailMessage(
                subject=subject,
                body=message,
                from_email=settings.DEFAULT_FROM_EMAIL,
                to=[recipient]
            )
            
            # Attach files
            for file_path in file_paths:
                email.attach_file(file_path)
            
            email.send()
            
            self.stdout.write(
                self.style.SUCCESS(f'Email sent successfully to {recipient}')
            )
        
        except Exception as e:
            self.stdout.write(
                self.style.ERROR(f'Failed to send email: {str(e)}')
            )